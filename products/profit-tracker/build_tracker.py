"""Builds 'The Side Hustle Profit Tracker' — a sellable Excel / Google Sheets workbook."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment
from datetime import date

OUT = "Side_Hustle_Profit_Tracker.xlsx"
ROWS = 500  # pre-formatted entry rows per log

INK, SAGE, SAGE_LT, GOLD = "1F2A2E", "5E8C7A", "E6EFEB", "D9A441"
F = "Arial"
H1 = Font(name=F, size=20, bold=True, color="FFFFFF")
H2 = Font(name=F, size=12, bold=True, color=INK)
HDR = Font(name=F, size=10, bold=True, color="FFFFFF")
BODY = Font(name=F, size=10, color=INK)
INPUT = Font(name=F, size=10, color="0000FF")
BOLD = Font(name=F, size=10, bold=True, color=INK)
MUTED = Font(name=F, size=9, italic=True, color="6B7A80")
FILL_SAGE = PatternFill("solid", fgColor=SAGE)
FILL_INK = PatternFill("solid", fgColor=INK)
FILL_LT = PatternFill("solid", fgColor=SAGE_LT)
FILL_YEL = PatternFill("solid", fgColor="FFF4CC")
THIN = Side(style="thin", color="C9D3CF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
USD = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = "mm/dd/yyyy"

wb = Workbook()


def banner(ws, title, sub, width):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws["A1"] = title
    ws["A1"].font = H1
    ws["A2"] = sub
    ws["A2"].font = Font(name=F, size=10, color="FFFFFF")
    for r in (1, 2):
        for c in range(1, width + 1):
            ws.cell(r, c).fill = FILL_SAGE
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 20
    for c in range(1, width + 1):
        ws.cell(3, c).fill = PatternFill("solid", fgColor=GOLD)
    ws.row_dimensions[3].height = 4


def log_sheet(name, title, sub, cols, example, formulas, widths):
    """cols: list of (header, kind) where kind in input/formula. Header row 5, data from row 6."""
    ws = wb.create_sheet(name)
    banner(ws, title, sub, len(cols))
    for i, (h, kind) in enumerate(cols, 1):
        c = ws.cell(5, i, h)
        c.font = HDR
        c.fill = FILL_INK if kind == "input" else PatternFill("solid", fgColor=SAGE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = widths[i - 1]
    ws.row_dimensions[5].height = 30
    ws["A4"] = "Blue columns (dark header) = you type.  Green columns = automatic, don't type here."
    ws["A4"].font = MUTED
    for r in range(6, 6 + ROWS):
        for i, (h, kind) in enumerate(cols, 1):
            c = ws.cell(r, i)
            c.border = BOX
            c.font = INPUT if kind == "input" else BODY
            if r % 2 == 1:
                c.fill = FILL_LT
            if i in formulas:
                c.value = formulas[i].format(r=r)
    for i, v in enumerate(example, 1):
        if v is not None and i not in formulas:
            ws.cell(6, i, v)
    ws.freeze_panes = "A6"
    return ws


# ---------------- Settings ----------------
st = wb.active
st.title = "Settings"
banner(st, "SETTINGS", "Set these once. Everything else updates on its own.", 4)
st.column_dimensions["A"].width = 34
st.column_dimensions["B"].width = 16
st.column_dimensions["C"].width = 4
st.column_dimensions["D"].width = 60
settings = [
    ("Tracking year", date.today().year, "0", "The year the Dashboard shows. Change it each January."),
    ("Tax set-aside rate", 0.25, PCT, "Share of profit to save for taxes. 25-30% is a common rule of thumb; ask a tax pro for your number."),
    ("Mileage rate ($ per mile)", 0.70, "$0.000", "IRS standard business mileage rate. 0.70 was the 2025 rate. Check irs.gov each year and update."),
    ("Monthly profit goal", 1000, USD, "Your target. The Dashboard shows how close you are each month."),
]
for i, (lab, val, fmt, note) in enumerate(settings):
    r = 5 + i
    st.cell(r, 1, lab).font = BOLD
    c = st.cell(r, 2, val)
    c.font = INPUT
    c.fill = FILL_YEL
    c.number_format = fmt
    c.border = BOX
    st.cell(r, 4, note).font = MUTED
    st.cell(r, 4).alignment = Alignment(wrap_text=True)
    st.row_dimensions[r].height = 28

st["A11"] = "INCOME SOURCES (edit to match your hustles)"
st["A11"].font = H2
sources = ["Marketplace flips", "eBay", "Etsy", "Shopify store", "DoorDash / Uber", "Freelance / services",
           "Content / creator payouts", "Affiliate", "Other"]
st["D11"] = "EXPENSE CATEGORIES"
st["D11"].font = H2
cats = ["Inventory / supplies", "Gas", "Phone & internet", "Software & apps", "Advertising", "Shipping & packaging",
        "Tools & equipment", "Fees", "Education", "Personal", "Other"]
for i in range(15):
    a = st.cell(12 + i, 1, sources[i] if i < len(sources) else None)
    a.font = INPUT
    a.border = BOX
    a.fill = FILL_YEL
    d = st.cell(12 + i, 4, cats[i] if i < len(cats) else None)
    d.font = INPUT
    d.border = BOX
    d.fill = FILL_YEL
st["A28"] = "Leave blank rows blank. Up to 15 of each. Yellow cells = edit."
st["A28"].font = MUTED

YEAR, TAX, MILE, GOAL = "Settings!$B$5", "Settings!$B$6", "Settings!$B$7", "Settings!$B$8"

# ---------------- Income ----------------
inc = log_sheet(
    "Income", "INCOME LOG", "Every dollar that comes in. Fees are what the platform kept.",
    [("Date", "input"), ("Source", "input"), ("Description", "input"), ("Gross ($)", "input"),
     ("Platform fees ($)", "input"), ("Net ($)", "formula"), ("Month", "formula"), ("Year", "formula")],
    [date(date.today().year, 1, 6), "Etsy", "3 printable planners", 14.97, 2.31],
    {6: '=IF(D{r}="","",D{r}-N(E{r}))', 7: '=IF(A{r}="","",MONTH(A{r}))', 8: '=IF(A{r}="","",YEAR(A{r}))'},
    [13, 22, 34, 13, 15, 13, 9, 9])
for r in range(6, 6 + ROWS):
    inc.cell(r, 1).number_format = DATE
    for col in (4, 5, 6):
        inc.cell(r, col).number_format = USD

# ---------------- Expenses ----------------
exp = log_sheet(
    "Expenses", "EXPENSE LOG", "Mark Business = Yes for anything you spent to run your hustle.",
    [("Date", "input"), ("Category", "input"), ("Description", "input"), ("Amount ($)", "input"),
     ("Business?", "input"), ("Month", "formula"), ("Year", "formula")],
    [date(date.today().year, 1, 8), "Shipping & packaging", "Poly mailers, 100 pack", 11.99, "Yes"],
    {6: '=IF(A{r}="","",MONTH(A{r}))', 7: '=IF(A{r}="","",YEAR(A{r}))'},
    [13, 24, 34, 13, 11, 9, 9])
for r in range(6, 6 + ROWS):
    exp.cell(r, 1).number_format = DATE
    exp.cell(r, 4).number_format = USD

# ---------------- Flips ----------------
fl = log_sheet(
    "Flips", "FLIP TRACKER", "One row per item. Profit, ROI and days-to-sell fill in when you add the sold price.",
    [("Item", "input"), ("Bought on", "input"), ("Buy price ($)", "input"), ("Repair / clean ($)", "input"),
     ("Listed at ($)", "input"), ("Sold on", "input"), ("Sold for ($)", "input"), ("Fees + shipping ($)", "input"),
     ("Profit ($)", "formula"), ("ROI", "formula"), ("Days to sell", "formula"), ("Status", "formula"),
     ("Month sold", "formula"), ("Year sold", "formula")],
    ["Dresser, solid wood (repainted)", date(date.today().year, 1, 3), 25, 18, 140,
     date(date.today().year, 1, 12), 125, 0],
    {9: '=IF(G{r}="","",G{r}-N(C{r})-N(D{r})-N(H{r}))',
     10: '=IF(I{r}="","",IF(N(C{r})+N(D{r})=0,"",I{r}/(N(C{r})+N(D{r}))))',
     11: '=IF(OR(B{r}="",F{r}=""),"",F{r}-B{r})',
     12: '=IF(A{r}="","",IF(G{r}<>"","Sold",IF(E{r}<>"","Listed","In stock")))',
     13: '=IF(F{r}="","",MONTH(F{r}))', 14: '=IF(F{r}="","",YEAR(F{r}))'},
    [32, 12, 12, 13, 12, 12, 12, 13, 12, 9, 10, 10, 8, 8])
for r in range(6, 6 + ROWS):
    for col in (2, 6):
        fl.cell(r, col).number_format = DATE
    for col in (3, 4, 5, 7, 8, 9):
        fl.cell(r, col).number_format = USD
    fl.cell(r, 10).number_format = PCT
    fl.cell(r, 11).number_format = "0"
fl.conditional_formatting.add("I6:I505", CellIsRule(operator="lessThan", formula=["0"],
                              font=Font(name=F, color="C0392B", bold=True)))

# ---------------- Subscriptions ----------------
sub = log_sheet(
    "Subscriptions", "SUBSCRIPTION AUDIT", "List every recurring charge. Mark the ones to cancel, then go cancel them.",
    [("Service", "input"), ("Category", "input"), ("Cost ($)", "input"), ("Billed", "input"),
     ("Next renewal", "input"), ("Keep / Cancel", "input"), ("Monthly cost ($)", "formula"),
     ("Yearly cost ($)", "formula"), ("Saved if cancelled ($/yr)", "formula")],
    ["Streaming service", "Personal", 15.49, "Monthly", date(date.today().year, 2, 1), "Cancel"],
    {7: '=IF(C{r}="","",IF(D{r}="Yearly",C{r}/12,IF(D{r}="Weekly",C{r}*52/12,IF(D{r}="Quarterly",C{r}/3,C{r}))))',
     8: '=IF(G{r}="","",G{r}*12)',
     9: '=IF(H{r}="","",IF(F{r}="Cancel",H{r},0))'},
    [26, 20, 11, 12, 13, 13, 14, 14, 16])
for r in range(6, 6 + ROWS):
    sub.cell(r, 5).number_format = DATE
    for col in (3, 7, 8, 9):
        sub.cell(r, col).number_format = USD
sub.conditional_formatting.add("A6:I505", FormulaRule(formula=['$F6="Cancel"'],
                               font=Font(name=F, color="C0392B", strike=True)))

# ---------------- Mileage ----------------
mi = log_sheet(
    "Mileage", "MILEAGE LOG", "Business trips only: sourcing runs, deliveries, post office, client jobs.",
    [("Date", "input"), ("Purpose", "input"), ("Start odometer", "input"), ("End odometer", "input"),
     ("Miles", "formula"), ("Deduction ($)", "formula"), ("Month", "formula"), ("Year", "formula")],
    [date(date.today().year, 1, 3), "Pick up dresser (flip)", 48210, 48232],
    {5: '=IF(OR(C{r}="",D{r}=""),"",D{r}-C{r})', 6: '=IF(E{r}="","",E{r}*' + MILE + ')',
     7: '=IF(A{r}="","",MONTH(A{r}))', 8: '=IF(A{r}="","",YEAR(A{r}))'},
    [13, 34, 15, 15, 10, 14, 9, 9])
for r in range(6, 6 + ROWS):
    mi.cell(r, 1).number_format = DATE
    mi.cell(r, 3).number_format = "#,##0"
    mi.cell(r, 4).number_format = "#,##0"
    mi.cell(r, 6).number_format = USD

# ---------------- Dropdowns ----------------
def dv(ws, formula, rng):
    v = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(v)
    v.add(rng)

dv(inc, "=Settings!$A$12:$A$26", "B6:B505")
dv(exp, "=Settings!$D$12:$D$26", "B6:B505")
dv(exp, '"Yes,No"', "E6:E505")
dv(sub, '"Business,Personal"', "B6:B505")
dv(sub, '"Weekly,Monthly,Quarterly,Yearly"', "D6:D505")
dv(sub, '"Keep,Cancel,Not sure"', "F6:F505")

# ---------------- Dashboard ----------------
db = wb.create_sheet("Dashboard", 0)
banner(db, "SIDE HUSTLE PROFIT TRACKER", "Your money at a glance. This page fills itself in from the other tabs.", 9)
widths = [14, 15, 15, 15, 15, 15, 15, 15, 15]
for i, w in enumerate(widths, 1):
    db.column_dimensions[chr(64 + i)].width = w

# KPI cards (row 5 label, row 6 value)
kpis = [
    ("A", "B", "PROFIT THIS YEAR", "=I20", USD),
    ("C", "D", "SET ASIDE FOR TAXES", "=MAX(0,I20)*" + TAX, USD),
    ("E", "F", "SUBSCRIPTIONS / MONTH", "=SUM(Subscriptions!G6:G505)", USD),
    ("G", "H", "AVG FLIP ROI", '=IFERROR(SUMIFS(Flips!I6:I505,Flips!N6:N505,' + YEAR + ')/(SUMIFS(Flips!C6:C505,Flips!N6:N505,' + YEAR + ')+SUMIFS(Flips!D6:D505,Flips!N6:N505,' + YEAR + ')),0)', PCT),
]
for a, b, lab, f, fmt in kpis:
    db.merge_cells(f"{a}5:{b}5")
    db.merge_cells(f"{a}6:{b}6")
    db[f"{a}5"] = lab
    db[f"{a}5"].font = Font(name=F, size=9, bold=True, color="6B7A80")
    db[f"{a}5"].alignment = Alignment(horizontal="center")
    db[f"{a}6"] = f
    db[f"{a}6"].font = Font(name=F, size=18, bold=True, color=INK)
    db[f"{a}6"].alignment = Alignment(horizontal="center")
    db[f"{a}6"].number_format = fmt
    for col in (a, b):
        for r in (5, 6):
            db[f"{col}{r}"].fill = FILL_LT
db.row_dimensions[6].height = 30
db["I5"] = "Year:"
db["I5"].font = MUTED
db["I6"] = "=" + YEAR
db["I6"].font = Font(name=F, size=14, bold=True, color=INK)
db["I6"].alignment = Alignment(horizontal="center")

heads = ["Month", "Income (net)", "Flip profit", "Business expenses", "Mileage deduction",
         "Profit", "Tax set-aside", "Goal progress", "Year-to-date profit"]
for i, h in enumerate(heads, 1):
    c = db.cell(7, i, h)
    c.font = HDR
    c.fill = FILL_INK
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BOX
db.row_dimensions[7].height = 30
months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
for m in range(1, 13):
    r = 7 + m
    db.cell(r, 1, months[m - 1]).font = BOLD
    db.cell(r, 2, f"=SUMIFS(Income!$F$6:$F$505,Income!$G$6:$G$505,{m},Income!$H$6:$H$505,{YEAR})")
    db.cell(r, 3, f"=SUMIFS(Flips!$I$6:$I$505,Flips!$M$6:$M$505,{m},Flips!$N$6:$N$505,{YEAR})")
    db.cell(r, 4, f'=SUMIFS(Expenses!$D$6:$D$505,Expenses!$F$6:$F$505,{m},Expenses!$G$6:$G$505,{YEAR},Expenses!$E$6:$E$505,"Yes")')
    db.cell(r, 5, f"=SUMIFS(Mileage!$F$6:$F$505,Mileage!$G$6:$G$505,{m},Mileage!$H$6:$H$505,{YEAR})")
    db.cell(r, 6, f"=B{r}+C{r}-D{r}-E{r}")
    db.cell(r, 7, f"=MAX(0,F{r})*{TAX}")
    db.cell(r, 8, f"=IF({GOAL}=0,0,F{r}/{GOAL})")
    db.cell(r, 9, f"=SUM($F$8:F{r})")
    for c in range(1, 10):
        cell = db.cell(r, c)
        cell.border = BOX
        if c > 1:
            cell.font = BODY
            cell.number_format = PCT if c == 8 else USD
        if m % 2 == 0:
            cell.fill = FILL_LT
r = 20
db.cell(r, 1, "TOTAL").font = HDR
for c in range(2, 8):
    if c == 8:
        continue
    col = chr(64 + c)
    db.cell(r, c, f"=SUM({col}8:{col}19)")
db.cell(r, 8, f"=IF({GOAL}=0,0,F20/({GOAL}*12))")
db.cell(r, 9, "=F20")
for c in range(1, 10):
    cell = db.cell(r, c)
    cell.fill = FILL_SAGE
    cell.border = BOX
    if c > 1:
        cell.font = HDR
        cell.number_format = PCT if c == 8 else USD
db.conditional_formatting.add("H8:H19", CellIsRule(operator="greaterThanOrEqual", formula=["1"],
                              fill=PatternFill("solid", fgColor="CDEBD9"), font=Font(name=F, bold=True, color="1E7B4B")))
db.conditional_formatting.add("F8:F20", CellIsRule(operator="lessThan", formula=["0"],
                              font=Font(name=F, bold=True, color="C0392B")))

db["A22"] = "FLIP STATS"
db["A22"].font = H2
stats = [
    ("Items sold (all time)", '=COUNTIF(Flips!L6:L505,"Sold")', "0"),
    ("Items listed, not sold", '=COUNTIF(Flips!L6:L505,"Listed")', "0"),
    ("Items in stock, not listed", '=COUNTIF(Flips!L6:L505,"In stock")', "0"),
    ("Money tied up in unsold items", '=SUMIFS(Flips!C6:C505,Flips!L6:L505,"<>Sold")+SUMIFS(Flips!D6:D505,Flips!L6:L505,"<>Sold")', USD),
    ("Average days to sell", '=IFERROR(AVERAGE(Flips!K6:K505),0)', "0.0"),
    ("Best flip profit", "=MAX(Flips!I6:I505)", USD),
]
for i, (lab, f, fmt) in enumerate(stats):
    rr = 23 + i
    db.merge_cells(f"A{rr}:C{rr}")
    db[f"A{rr}"] = lab
    db[f"A{rr}"].font = BODY
    db[f"D{rr}"] = f
    db[f"D{rr}"].font = BOLD
    db[f"D{rr}"].number_format = fmt
db["F22"] = "SUBSCRIPTIONS"
db["F22"].font = H2
substats = [
    ("Monthly total", "=SUM(Subscriptions!G6:G505)", USD),
    ("Yearly total", "=SUM(Subscriptions!H6:H505)", USD),
    ("Yearly savings marked Cancel", "=SUM(Subscriptions!I6:I505)", USD),
]
for i, (lab, f, fmt) in enumerate(substats):
    rr = 23 + i
    db.merge_cells(f"F{rr}:H{rr}")
    db[f"F{rr}"] = lab
    db[f"F{rr}"].font = BODY
    db[f"I{rr}"] = f
    db[f"I{rr}"].font = BOLD
    db[f"I{rr}"].number_format = fmt
db["A30"] = ("Profit = net income + flip profit - business expenses - mileage deduction. "
             "Tax set-aside is an estimate for saving, not tax advice.")
db["A30"].font = MUTED
db.freeze_panes = "A8"

# ---------------- Start Here ----------------
sh = wb.create_sheet("Start Here", 0)
banner(sh, "START HERE", "5 minutes to set up. 10 minutes a week to keep it running.", 2)
sh.column_dimensions["A"].width = 6
sh.column_dimensions["B"].width = 100
steps = [
    ("SETUP (one time)", None),
    ("1", "Go to the Settings tab. Set your year, tax set-aside rate, mileage rate and monthly profit goal (yellow cells)."),
    ("2", "Edit the income sources and expense categories in Settings to match your hustles. The dropdowns update automatically."),
    ("3", "Delete the example row (row 6) on each log tab once you've seen how it works. Clear the blue cells only."),
    ("WEEKLY (10 minutes)", None),
    ("4", "Income: add every payout. Put what the platform kept in Platform fees."),
    ("5", "Expenses: add what you spent. Mark Business = Yes for hustle costs."),
    ("6", "Flips: add items when you buy them. Add Sold on + Sold for when they sell; profit and ROI calculate themselves."),
    ("7", "Mileage: log business trips with your odometer readings."),
    ("8", "Check the Dashboard. Move the Tax set-aside amount into a separate savings account."),
    ("MONTHLY", None),
    ("9", "Subscriptions: review every recurring charge. Mark Cancel, then actually cancel it the same day."),
    ("COLOR KEY", None),
    ("", "Blue text / dark header = you type here.   Black text / green header = automatic, don't type here.   Yellow = settings."),
    ("GOOGLE SHEETS", None),
    ("", "Open Google Drive > New > File upload > choose this file > open it > File > Save as Google Sheets. Everything works there too."),
    ("NOTE", None),
    ("", "This tracker helps you organize records. It isn't tax or legal advice. Keep your receipts and talk to a tax pro about your situation."),
]
r = 5
for a, b in steps:
    if b is None:
        sh.cell(r, 1, a).font = H2
        r += 1
        continue
    sh.cell(r, 1, a).font = Font(name=F, size=11, bold=True, color=SAGE)
    sh.cell(r, 2, b).font = Font(name=F, size=11, color=INK)
    sh.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
    sh.row_dimensions[r].height = 30
    r += 1
sh["B" + str(r + 1)] = "The Side Hustle Profit Tracker  |  For personal use only. Not for resale."
sh["B" + str(r + 1)].font = MUTED

# Tab colors
for name, color in [("Start Here", GOLD), ("Dashboard", SAGE), ("Settings", "999999")]:
    wb[name].sheet_properties.tabColor = color
wb.active = 0
wb.save(OUT)
print("saved", OUT)
